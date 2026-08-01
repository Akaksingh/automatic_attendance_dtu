"""Background tasks for account provisioning and face-photo processing."""
import json
import csv
import io
import secrets
import string
from datetime import UTC, datetime
from pathlib import Path
from uuid import UUID
from zipfile import ZipFile

import numpy as np
from openpyxl import load_workbook

from .celery_app import celery_app
from .config import settings
from .database import SessionLocal
from .face.engine import cosine_similarity, detect_faces, embed_single_face
from .email_service import send_welcome_email
from .models import AttendanceRecord, AttendanceStatus, ClassSession, Group, GroupMembership, ImportJob, ImportJobStatus, PhotoJobStatus, PhotoProcessingJob, RefreshToken, SessionPhoto, SessionState, Student, StudentPhoto, User, UserRole
from .security import hash_password
from .storage import storage


@celery_app.task(name="app.tasks.ping")
def ping() -> dict[str, str]:
    """Minimal task that confirms the worker can consume a Celery job."""
    return {"status": "ok"}


@celery_app.task(bind=True, name="app.tasks.send_welcome_email", max_retries=3)
def send_welcome_email_task(self, email: str, full_name: str, password: str) -> dict[str, str]:
    try:
        send_welcome_email(email=email, full_name=full_name, password=password)
        return {"status": "sent", "email": email}
    except Exception as exc:
        raise self.retry(exc=exc, countdown=60)


def _temporary_password() -> str:
    alphabet = string.ascii_letters + string.digits
    return "".join(secrets.choice(alphabet) for _ in range(10))


def _canonical(value: object) -> str:
    return "".join(character for character in str(value or "").lower() if character.isalnum())


def _error_file(job: ImportJob, errors: list[tuple[int, str]], db) -> None:
    if not errors:
        return
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["row", "error"])
    writer.writerows(errors)
    key = f"imports/{job.id}/errors.csv"
    storage.put_bytes(key, output.getvalue().encode(), "text/csv")
    job.error_file_key = key


def _finish_import(job: ImportJob, db, errors: list[tuple[int, str]]) -> None:
    _error_file(job, errors, db)
    job.failed = len(errors)
    job.status = ImportJobStatus.completed
    job.completed_at = datetime.now(UTC)
    storage.delete(job.spreadsheet_key)
    if job.images_zip_key:
        storage.delete(job.images_zip_key)
    db.commit()


def _match_faces(face_embeddings: list[np.ndarray], references: dict[UUID, list[np.ndarray]], threshold: float) -> dict[UUID, float]:
    """Greedily assign each detected face and student at most once.

    This mirrors the proven proof-of-concept matcher while still comparing a
    student against all of their approved reference embeddings.
    """
    if not face_embeddings or not references:
        return {}
    student_ids = list(references)
    scores = np.zeros((len(face_embeddings), len(student_ids)), dtype=np.float32)
    for face_index, face in enumerate(face_embeddings):
        for student_index, student_id in enumerate(student_ids):
            scores[face_index, student_index] = max(cosine_similarity(face, reference) for reference in references[student_id])
    face_indices, student_indices = np.unravel_index(np.argsort(scores, axis=None)[::-1], scores.shape)
    matched: dict[UUID, float] = {}
    used_faces: set[int] = set()
    used_students: set[int] = set()
    for face_index, student_index in zip(face_indices.tolist(), student_indices.tolist()):
        score = float(scores[face_index, student_index])
        if score < threshold:
            break
        if face_index in used_faces or student_index in used_students:
            continue
        used_faces.add(face_index)
        used_students.add(student_index)
        matched[student_ids[student_index]] = score
    return matched


@celery_app.task(bind=True, name="app.tasks.process_import_job", max_retries=2)
def process_import_job(self, import_job_id: UUID) -> dict[str, object]:
    """Process all spreadsheet imports off the request thread."""
    with SessionLocal() as db:
        job = db.get(ImportJob, import_job_id)
        if job is None:
            return {"status": "missing", "job_id": str(import_job_id)}
        if job.status == ImportJobStatus.completed:
            return {"status": "completed", "job_id": str(import_job_id)}
        job.status = ImportJobStatus.processing
        errors: list[tuple[int, str]] = []
        try:
            workbook = load_workbook(io.BytesIO(storage.get_bytes(job.spreadsheet_key)), read_only=True, data_only=True)
            rows = workbook.active.iter_rows(values_only=True)
            headers = [_canonical(value) for value in next(rows, ())]
            values = list(rows)
            job.total = len(values)
            db.commit()

            zip_entries: dict[str, tuple[str, bytes]] = {}
            if job.images_zip_key:
                with ZipFile(io.BytesIO(storage.get_bytes(job.images_zip_key))) as archive:
                    for info in archive.infolist():
                        if info.is_dir() or ".." in Path(info.filename).parts:
                            continue
                        normalized = info.filename.replace("\\", "/").lower()
                        zip_entries[normalized] = (Path(normalized).name, archive.read(info))

            if job.kind == "student_create":
                required = {"email", "fullname", "rollno"}
                if not required.issubset(headers):
                    raise ValueError("Excel must contain email, full name, and roll no columns")
                for row_number, row in enumerate(values, start=2):
                    photo_job = None
                    try:
                        data = {headers[index]: row[index] if index < len(row) else None for index in range(len(headers))}
                        email = str(data.get("email") or "").strip().lower()
                        full_name = str(data.get("fullname") or "").strip()
                        roll_no = str(data.get("rollno") or "").strip()
                        if db.query(User).filter(User.email == email).first() or db.query(Student).filter(Student.roll_no == roll_no).first():
                            raise ValueError("Email or roll number is already registered")
                        supplied_password = str(data.get("password") or "").strip()
                        password = supplied_password or _temporary_password()
                        user = User(email=email, full_name=full_name, password_hash=hash_password(password), role=UserRole.student)
                        db.add(user)
                        db.flush()
                        student = Student(user_id=user.id, roll_no=roll_no, full_name=full_name)
                        db.add(student)
                        # The photo job FK needs the generated student UUID
                        # before the job is inserted.
                        db.flush()
                        photo_name = str(data.get("photo") or "").strip().replace("\\", "/").lower()
                        if photo_name:
                            entry = zip_entries.get(photo_name) or next((value for key, value in zip_entries.items() if Path(key).name == Path(photo_name).name), None)
                            if entry is None:
                                raise ValueError("Photo is missing from images_zip")
                            file_name, content = entry
                            suffix = Path(file_name).suffix.lower()
                            content_type = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp"}.get(suffix)
                            if content_type is None or len(content) > settings.max_profile_photo_bytes:
                                raise ValueError("Photo must be a supported image within the size limit")
                            key = f"profile/pending/{student.id}/{secrets.token_urlsafe(24)}{suffix}"
                            storage.put_bytes(key, content, content_type)
                            photo_job = PhotoProcessingJob(student_id=student.id, storage_key=key, content_type=content_type)
                            db.add(photo_job)
                            db.flush()
                        db.commit()
                        if photo_job:
                            process_student_photo.delay(photo_job.id)
                        send_welcome_email_task.delay(email, full_name, password)
                        job.succeeded += 1
                    except Exception as exc:
                        db.rollback()
                        errors.append((row_number, str(exc)))
                    job.processed += 1
                    db.commit()
                    self.update_state(state="PROGRESS", meta={"processed": job.processed, "total": job.total, "failed": len(errors)})
            elif job.kind == "student_update":
                required = {"rollno", "newrollno", "fullname"}
                if not required.issubset(headers):
                    raise ValueError("Excel must contain RollNo, New RollNo, and Full Name columns")
                for row_number, row in enumerate(values, start=2):
                    try:
                        data = {headers[index]: row[index] if index < len(row) else None for index in range(len(headers))}
                        current = db.query(Student).filter(Student.roll_no == str(data.get("rollno") or "").strip()).first()
                        new_roll = str(data.get("newrollno") or "").strip()
                        full_name = str(data.get("fullname") or "").strip()
                        if current is None:
                            raise ValueError("Current roll number was not found")
                        if db.query(Student).filter(Student.roll_no == new_roll, Student.id != current.id).first():
                            raise ValueError("New roll number is already registered")
                        email = str(data.get("email") or current.user.email).strip().lower()
                        if db.query(User).filter(User.email == email, User.id != current.user_id).first():
                            raise ValueError("Email is already registered")
                        current.roll_no = new_roll
                        current.full_name = full_name
                        current.user.full_name = full_name
                        current.user.email = email
                        db.commit()
                        job.succeeded += 1
                    except Exception as exc:
                        db.rollback()
                        errors.append((row_number, str(exc)))
                    job.processed += 1
                    db.commit()
                    self.update_state(state="PROGRESS", meta={"processed": job.processed, "total": job.total, "failed": len(errors)})
            elif job.kind == "teacher_create":
                required = {"email", "fullname"}
                if not required.issubset(headers):
                    raise ValueError("Excel must contain email and full name columns")
                for row_number, row in enumerate(values, start=2):
                    try:
                        data = {headers[index]: row[index] if index < len(row) else None for index in range(len(headers))}
                        email = str(data.get("email") or "").strip().lower()
                        full_name = str(data.get("fullname") or "").strip()
                        if not email or not full_name or db.query(User).filter(User.email == email).first():
                            raise ValueError("Valid name/email are required and email must be unique")
                        password = str(data.get("password") or "").strip() or _temporary_password()
                        teacher = User(email=email, full_name=full_name, password_hash=hash_password(password), role=UserRole.teacher)
                        db.add(teacher)
                        db.commit()
                        send_welcome_email_task.delay(email, full_name, password)
                        job.succeeded += 1
                    except Exception as exc:
                        db.rollback()
                        errors.append((row_number, str(exc)))
                    job.processed += 1
                    db.commit()
                    self.update_state(state="PROGRESS", meta={"processed": job.processed, "total": job.total, "failed": len(errors)})
            elif job.kind == "group_roster":
                required = {"rollno", "fullname"}
                if not required.issubset(headers):
                    raise ValueError("Excel must contain rollNo and full name columns")
                group = db.get(Group, job.group_id)
                if group is None:
                    raise ValueError("Target group no longer exists")
                for row_number, row in enumerate(values, start=2):
                    try:
                        data = {headers[index]: row[index] if index < len(row) else None for index in range(len(headers))}
                        roll_no = str(data.get("rollno") or "").strip()
                        student = db.query(Student).filter(Student.roll_no == roll_no).first()
                        if student is None:
                            raise ValueError("Student roll number was not found")
                        if db.query(GroupMembership).filter_by(group_id=group.id, student_id=student.user_id).first():
                            raise ValueError("Student is already in this group")
                        db.add(GroupMembership(group_id=group.id, student_id=student.user_id))
                        db.commit()
                        job.succeeded += 1
                    except Exception as exc:
                        db.rollback()
                        errors.append((row_number, str(exc)))
                    job.processed += 1
                    db.commit()
                    self.update_state(state="PROGRESS", meta={"processed": job.processed, "total": job.total, "failed": len(errors)})
            else:
                raise ValueError(f"Unknown import type: {job.kind}")
            workbook.close()
            _finish_import(job, db, errors)
            return {"status": "completed", "job_id": str(job.id), "processed": job.processed, "failed": len(errors)}
        except Exception as exc:
            db.rollback()
            job = db.get(ImportJob, import_job_id)
            job.status = ImportJobStatus.failed
            job.error = str(exc)
            db.commit()
            raise self.retry(exc=exc, countdown=30)


@celery_app.task(bind=True, name="app.tasks.process_student_photo", max_retries=2)
def process_student_photo(self, job_id: UUID) -> dict[str, object]:
    """Verify one staged image and promote it to the student's photo history."""
    with SessionLocal() as db:
        job = db.get(PhotoProcessingJob, job_id)
        if job is None:
            return {"status": "missing", "job_id": job_id}
        if job.status in {PhotoJobStatus.approved, PhotoJobStatus.rejected}:
            return {"status": job.status.value, "job_id": job_id}
        job.status = PhotoJobStatus.processing
        db.commit()
        try:
            if not storage.exists(job.storage_key):
                raise ValueError("Staged photo file is missing")
            embedding = embed_single_face(storage.get_bytes(job.storage_key))
            student = db.get(Student, job.student_id)
            if student is None:
                raise ValueError("Student profile no longer exists")
            previous = db.query(StudentPhoto).filter(StudentPhoto.student_id == student.id).all()
            similarity = None
            if previous:
                similarity = max(cosine_similarity(embedding, np.asarray(json.loads(photo.embedding_json), dtype=np.float32)) for photo in previous)
                job.similarity_to_previous = similarity
                if similarity < settings.face_match_threshold:
                    job.status = PhotoJobStatus.rejected
                    job.error = f"Photo does not match the existing profile (similarity {similarity:.3f})"
                    job.completed_at = datetime.now(UTC)
                    db.commit()
                    storage.delete(job.storage_key)
                    return {"status": "rejected", "job_id": job_id, "similarity": similarity}
            for old_photo in previous:
                old_photo.is_profile = False
                if old_photo.storage_key:
                    storage.delete(old_photo.storage_key)
                    old_photo.storage_key = None
            photo = StudentPhoto(student_id=student.id, storage_key=job.storage_key, content_type=job.content_type, embedding_json=json.dumps(embedding.tolist()), similarity_to_previous=similarity, is_profile=True)
            db.add(photo)
            db.flush()
            student.profile_image_url = f"/api/students/photos/{photo.id}/image"
            student.user.profile_image_url = student.profile_image_url
            job.photo_id = photo.id
            job.status = PhotoJobStatus.approved
            job.completed_at = datetime.now(UTC)
            db.commit()
            return {"status": "approved", "job_id": job_id, "photo_id": photo.id, "similarity": similarity}
        except (ValueError, RuntimeError) as exc:
            db.rollback()
            job = db.get(PhotoProcessingJob, job_id)
            job.status = PhotoJobStatus.failed
            job.error = str(exc)
            job.completed_at = datetime.now(UTC)
            db.commit()
            storage.delete(job.storage_key)
            return {"status": "failed", "job_id": job_id, "error": str(exc)}
        except Exception as exc:
            db.rollback()
            job = db.get(PhotoProcessingJob, job_id)
            job.status = PhotoJobStatus.failed
            job.error = "Unexpected photo processing error"
            job.completed_at = datetime.now(UTC)
            db.commit()
            storage.delete(job.storage_key)
            raise self.retry(exc=exc, countdown=30)


@celery_app.task(bind=True, name="app.tasks.process_class_session", max_retries=2)
def process_class_session(self, session_id: UUID) -> dict[str, object]:
    """Run face matching for every classroom image and create reviewable records."""
    with SessionLocal() as db:
        session = db.get(ClassSession, session_id)
        if session is None:
            return {"status": "missing", "session_id": str(session_id)}
        session.state = SessionState.processing
        session.processing_total = len(session.photos)
        session.processing_completed = 0
        session.processing_message = "Preparing classroom images"
        db.commit()
        self.update_state(state="PROGRESS", meta={"completed": 0, "total": len(session.photos), "message": session.processing_message})
        try:
            memberships = db.query(GroupMembership).filter(GroupMembership.group_id == session.group_id).all()
            user_ids = [membership.student_id for membership in memberships]
            students = db.query(Student).filter(Student.user_id.in_(user_ids)).all() if user_ids else []
            references: dict[UUID, list[np.ndarray]] = {
                student.user_id: [np.asarray(json.loads(photo.embedding_json), dtype=np.float32) for photo in student.photos]
                for student in students
            }
            live_embeddings: list[np.ndarray] = []
            for photo in session.photos:
                if not storage.exists(photo.storage_key):
                    raise ValueError("Session image file is missing")
                faces = detect_faces(storage.get_bytes(photo.storage_key))
                live_embeddings.extend(face.embedding for face in faces)
                photo.processed_at = datetime.now(UTC)
                session.processing_completed += 1
                session.processing_message = f"Processed {session.processing_completed} of {session.processing_total} classroom images"
                db.commit()
                self.update_state(state="PROGRESS", meta={"completed": session.processing_completed, "total": session.processing_total, "message": session.processing_message})

            recognized = _match_faces(live_embeddings, references, settings.face_match_threshold)
            for user_id in user_ids:
                record = db.query(AttendanceRecord).filter_by(session_id=session.id, student_id=user_id).first()
                status = AttendanceStatus.present if user_id in recognized else AttendanceStatus.absent
                if record is None:
                    record = AttendanceRecord(session_id=session.id, student_id=user_id, status=status)
                    db.add(record)
                else:
                    record.status = status
                    record.marked_at = datetime.now(UTC)
            session.state = SessionState.review
            session.processing_message = "Attendance is ready for review"
            db.commit()
            return {"status": "review", "session_id": str(session_id), "recognized": len(recognized), "roster": len(user_ids)}
        except (ValueError, RuntimeError) as exc:
            db.rollback()
            session = db.get(ClassSession, session_id)
            session.state = SessionState.failed
            session.processing_message = str(exc)[:255]
            for photo in session.photos:
                photo.error = str(exc)
            db.commit()
            return {"status": "failed", "session_id": str(session_id), "error": str(exc)}
        except Exception as exc:
            db.rollback()
            session = db.get(ClassSession, session_id)
            session.state = SessionState.failed
            session.processing_message = "Unexpected attendance processing error"
            db.commit()
            raise self.retry(exc=exc, countdown=30)


@celery_app.task(name="app.tasks.expire_session_photo")
def expire_session_photo(photo_id: UUID) -> dict[str, object]:
    with SessionLocal() as db:
        photo = db.get(SessionPhoto, photo_id)
        if photo is None:
            return {"status": "missing", "photo_id": str(photo_id)}
        if photo.expires_at <= datetime.now(UTC):
            storage.delete(photo.storage_key)
            return {"status": "expired", "photo_id": str(photo_id)}
        return {"status": "not_due", "photo_id": str(photo_id)}


@celery_app.task(name="app.tasks.cleanup_expired_session_photos")
def cleanup_expired_session_photos() -> dict[str, int]:
    """Periodic retention backstop for images whose delayed task was lost."""
    now = datetime.now(UTC)
    removed = 0
    with SessionLocal() as db:
        photos = db.query(SessionPhoto).filter(SessionPhoto.expires_at <= now).all()
        for photo in photos:
            storage.delete(photo.storage_key)
            removed += 1
    return {"removed": removed}


@celery_app.task(name="app.tasks.cleanup_expired_refresh_tokens")
def cleanup_expired_refresh_tokens() -> dict[str, int]:
    with SessionLocal() as db:
        deleted = db.query(RefreshToken).filter(RefreshToken.expires_at <= datetime.now(UTC)).delete(synchronize_session=False)
        db.commit()
    return {"deleted": deleted}
