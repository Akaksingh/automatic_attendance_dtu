"""Admin-only account and bulk student provisioning."""
from io import BytesIO
from pathlib import Path
import secrets
from datetime import UTC, datetime
from typing import Annotated
from zipfile import BadZipFile, ZipFile
from uuid import UUID

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile, status
from openpyxl import load_workbook
from pydantic import EmailStr, TypeAdapter, ValidationError
from sqlalchemy import or_
from sqlalchemy.orm import Session

from ..config import settings
from ..database import get_db
from ..models import AuditLog, ImportJob, PhotoProcessingJob, RefreshToken, Student, User, UserRole
from ..schemas import (
    AdminStudentUpdate,
    BulkStudentImportOut,
    BulkStudentResult,
    PhotoJobOut,
    StudentOut,
    StudentPage,
    ImportJobOut,
    StudentProvisionOut,
    TeacherCreate,
    UserPage,
    UserOut,
)
from ..security import hash_password, require_roles
from ..tasks import process_import_job, process_student_photo, send_welcome_email_task
from ..storage import storage

router = APIRouter(prefix="/api/admin", tags=["admin"])
DbSession = Annotated[Session, Depends(get_db)]
email_adapter = TypeAdapter(EmailStr)


def _temporary_password() -> str:
    alphabet = "ABCDEFGHJKLMNPQRSTUVWXYZabcdefghijkmnopqrstuvwxyz23456789"
    return "".join(secrets.choice(alphabet) for _ in range(10))


def _stage_photo(student_id: UUID, file_name: str, content_type: str, content: bytes, db: Session) -> PhotoProcessingJob:
    allowed = {"image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp"}
    if content_type not in allowed:
        raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="Use JPEG, PNG, or WebP")
    if len(content) > settings.max_profile_photo_bytes:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Profile photo is too large")
    suffix = Path(file_name).suffix.lower() or allowed[content_type]
    if suffix not in {".jpg", ".jpeg", ".png", ".webp"}:
        suffix = allowed[content_type]
    key = f"profile/pending/{student_id}/{secrets.token_urlsafe(24)}{suffix}"
    storage.put_bytes(key, content, content_type)
    job = PhotoProcessingJob(student_id=student_id, storage_key=key, content_type=content_type)
    db.add(job)
    return job


def _create_student(db: Session, *, email: str, full_name: str, roll_no: str, password: str) -> Student:
    email = str(email).strip().lower()
    roll_no = roll_no.strip()
    full_name = full_name.strip()
    try:
        email = str(email_adapter.validate_python(email))
    except ValidationError as exc:
        raise ValueError("Email is invalid") from exc
    if len(full_name) < 2 or not roll_no:
        raise ValueError("Name and roll number are required")
    if db.query(User).filter(User.email == email).first():
        raise ValueError("Email is already registered")
    if db.query(Student).filter(Student.roll_no == roll_no).first():
        raise ValueError("Roll number is already registered")
    user = User(email=email, full_name=full_name, password_hash=hash_password(password), role=UserRole.student)
    db.add(user)
    db.flush()
    student = Student(user_id=user.id, roll_no=roll_no, full_name=full_name)
    db.add(student)
    db.flush()
    return student


def _student_out(student: Student) -> StudentOut:
    return StudentOut(
        id=student.id,
        user_id=student.user_id,
        roll_no=student.roll_no,
        full_name=student.full_name,
        profile_image_url=student.profile_image_url,
        photo_count=len(student.photos),
        photos=[],
    )


@router.post("/teachers", response_model=UserOut, status_code=status.HTTP_201_CREATED)
def create_teacher(
    payload: TeacherCreate,
    db: DbSession,
    user: User = Depends(require_roles(UserRole.admin)),
) -> User:
    email = payload.email.lower()
    if db.query(User).filter(User.email == email).first():
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Email is already registered")
    teacher = User(email=email, full_name=payload.full_name, password_hash=hash_password(payload.password), role=UserRole.teacher)
    db.add(teacher)
    db.commit()
    db.refresh(teacher)
    return teacher


@router.post("/users/{user_id}/reset-password", response_model=UserOut)
def reset_user_password(
    user_id: UUID,
    db: DbSession,
    admin: User = Depends(require_roles(UserRole.admin)),
) -> User:
    target = db.get(User, user_id)
    if target is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")
    temporary_password = _temporary_password()
    target.password_hash = hash_password(temporary_password)
    now = datetime.now(UTC)
    db.query(RefreshToken).filter(
        RefreshToken.user_id == target.id,
        RefreshToken.revoked_at.is_(None),
    ).update({RefreshToken.revoked_at: now}, synchronize_session=False)
    db.add(AuditLog(actor_id=admin.id, action="user.password_reset", entity_type="user", entity_id=target.id))
    db.commit()
    send_welcome_email_task.delay(target.email, target.full_name, temporary_password)
    db.refresh(target)
    return target


@router.post("/users/{user_id}/disable", response_model=UserOut)
def disable_user(
    user_id: UUID,
    db: DbSession,
    admin: User = Depends(require_roles(UserRole.admin)),
) -> User:
    target = db.get(User, user_id)
    if target is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")
    if target.id == admin.id:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="You cannot disable your own account")
    target.is_active = False
    now = datetime.now(UTC)
    db.query(RefreshToken).filter(
        RefreshToken.user_id == target.id,
        RefreshToken.revoked_at.is_(None),
    ).update({RefreshToken.revoked_at: now}, synchronize_session=False)
    db.add(AuditLog(actor_id=admin.id, action="user.disabled", entity_type="user", entity_id=target.id))
    db.commit()
    db.refresh(target)
    return target


@router.post("/users/{user_id}/enable", response_model=UserOut)
def enable_user(
    user_id: UUID,
    db: DbSession,
    admin: User = Depends(require_roles(UserRole.admin)),
) -> User:
    target = db.get(User, user_id)
    if target is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")
    target.is_active = True
    db.add(AuditLog(actor_id=admin.id, action="user.enabled", entity_type="user", entity_id=target.id))
    db.commit()
    db.refresh(target)
    return target


@router.get("/users", response_model=UserPage)
def list_users(
    db: DbSession,
    _: User = Depends(require_roles(UserRole.admin)),
    page: int = Query(1, ge=1),
    page_size: int = Query(25, ge=1, le=100),
    role: UserRole | None = Query(None),
    search: str | None = Query(None, max_length=100),
) -> UserPage:
    query = db.query(User).outerjoin(Student, Student.user_id == User.id).order_by(User.created_at.desc())
    if role is not None:
        query = query.filter(User.role == role)
    if search and search.strip():
        term = f"%{search.strip()}%"
        query = query.filter(or_(User.email.ilike(term), User.full_name.ilike(term), Student.roll_no.ilike(term)))
    total = query.count()
    users = query.offset((page - 1) * page_size).limit(page_size).all()
    items = []
    for item in users:
        output = UserOut.model_validate(item)
        student = db.query(Student).filter(Student.user_id == item.id).first()
        output.roll_no = student.roll_no if student else None
        items.append(output)
    return UserPage(items=items, meta={"page": page, "page_size": page_size, "total": total, "pages": (total + page_size - 1) // page_size})


@router.patch("/students/{user_id}", response_model=UserOut)
def update_student_details(
    user_id: UUID,
    payload: AdminStudentUpdate,
    db: DbSession,
    admin: User = Depends(require_roles(UserRole.admin)),
) -> User:
    target = db.get(User, user_id)
    student = db.query(Student).filter(Student.user_id == user_id).first()
    if target is None or student is None or target.role != UserRole.student:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Student not found")
    email = str(payload.email).strip().lower()
    duplicate_email = db.query(User).filter(User.email == email, User.id != target.id).first()
    if duplicate_email:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Email is already registered")
    duplicate_roll = db.query(Student).filter(Student.roll_no == payload.roll_no.strip(), Student.id != student.id).first()
    if duplicate_roll:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Roll number is already registered")
    target.email = email
    target.full_name = payload.full_name.strip()
    student.full_name = payload.full_name.strip()
    student.roll_no = payload.roll_no.strip()
    db.add(AuditLog(actor_id=admin.id, action="student.details_updated", entity_type="user", entity_id=target.id))
    db.commit()
    db.refresh(target)
    output = UserOut.model_validate(target)
    output.roll_no = student.roll_no
    return output


@router.get("/students", response_model=StudentPage)
def list_students(
    db: DbSession,
    _: User = Depends(require_roles(UserRole.admin)),
    page: int = Query(1, ge=1),
    page_size: int = Query(25, ge=1, le=100),
) -> StudentPage:
    query = db.query(Student).order_by(Student.roll_no.asc())
    total = query.count()
    students = query.offset((page - 1) * page_size).limit(page_size).all()
    return StudentPage(items=[_student_out(student) for student in students], meta={"page": page, "page_size": page_size, "total": total, "pages": (total + page_size - 1) // page_size})


@router.post("/students", response_model=StudentProvisionOut, status_code=status.HTTP_201_CREATED)
async def create_student(
    db: DbSession,
    email: str = Form(...),
    full_name: str = Form(...),
    roll_no: str = Form(...),
    password: str | None = Form(None),
    file: UploadFile | None = File(None),
    _: User = Depends(require_roles(UserRole.admin)),
) -> StudentProvisionOut:
    temporary_password = _temporary_password()
    job = None
    try:
        student = _create_student(db, email=email, full_name=full_name, roll_no=roll_no, password=temporary_password)
        if file is not None:
            job = _stage_photo(student.id, file.filename or "profile.jpg", file.content_type or "", await file.read(), db)
        db.commit()
    except (ValueError, HTTPException) as exc:
        db.rollback()
        if isinstance(exc, HTTPException):
            raise
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc)) from exc
    if job:
        process_student_photo.delay(job.id)
    send_welcome_email_task.delay(student.user.email, student.full_name, temporary_password)
    db.refresh(student)
    return StudentProvisionOut(student=_student_out(student), photo_job=PhotoJobOut.model_validate(job) if job else None, temporary_password=temporary_password)


@router.post("/students/bulk-legacy", response_model=BulkStudentImportOut, status_code=status.HTTP_202_ACCEPTED, include_in_schema=False)
async def bulk_create_students(
    db: DbSession,
    spreadsheet: UploadFile = File(...),
    images_zip: UploadFile | None = File(None),
    _: User = Depends(require_roles(UserRole.admin)),
) -> BulkStudentImportOut:
    if not (spreadsheet.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="Upload an .xlsx or .xlsm Excel file")
    image_entries: dict[str, tuple[str, bytes]] = {}
    if images_zip is not None:
        try:
            zip_content = await images_zip.read(settings.max_bulk_zip_bytes + 1)
            if len(zip_content) > settings.max_bulk_zip_bytes:
                raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="images_zip is too large")
            with ZipFile(BytesIO(zip_content)) as archive:
                if len(archive.infolist()) > settings.max_bulk_zip_files:
                    raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="images_zip contains too many files")
                for info in archive.infolist():
                    if info.is_dir():
                        continue
                    normalized = info.filename.replace("\\", "/").lstrip("./")
                    if ".." in Path(normalized).parts:
                        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Image ZIP contains an unsafe path")
                    if info.file_size > settings.max_profile_photo_bytes:
                        continue
                    image_entries[normalized.lower()] = (Path(normalized).name, archive.read(info))
        except BadZipFile as exc:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="images_zip is not a valid ZIP archive") from exc
    try:
        workbook = load_workbook(BytesIO(await spreadsheet.read()), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Could not read Excel file") from exc
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip().lower() for value in next(rows, ())]
    required = {"email", "full_name", "roll_no"}
    if not required.issubset(headers):
        workbook.close()
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Excel must contain email, full_name, and roll_no columns")
    results: list[BulkStudentResult] = []
    for row_number, values in enumerate(rows, start=2):
        data = {headers[index]: values[index] if index < len(values) else None for index in range(len(headers))}
        email = str(data.get("email") or "").strip().lower()
        roll_no = str(data.get("roll_no") or "").strip()
        full_name = str(data.get("full_name") or "").strip()
        supplied_password = str(data.get("password") or "").strip()
        password = supplied_password or _temporary_password()
        job = None
        try:
            student = _create_student(db, email=email, full_name=full_name, roll_no=roll_no, password=password)
            photo_name = str(data.get("photo") or "").strip()
            if photo_name:
                normalized_name = photo_name.replace("\\", "/").lstrip("./").lower()
                entry = image_entries.get(normalized_name)
                if entry is None:
                    matches = [value for key, value in image_entries.items() if Path(key).name == Path(normalized_name).name]
                    entry = matches[0] if len(matches) == 1 else None
                if entry is None:
                    raise ValueError("Photo is missing from images_zip")
                file_name, content = entry
                suffix = Path(file_name).suffix.lower()
                content_type = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp"}.get(suffix)
                if content_type is None:
                    raise ValueError("Photo must be JPEG, PNG, or WebP")
                job = _stage_photo(student.id, file_name, content_type, content, db)
            db.commit()
            if job:
                process_student_photo.delay(job.id)
            results.append(BulkStudentResult(row=row_number, email=email, roll_no=roll_no, status="created", student_id=student.id, photo_job_id=job.id if job else None, temporary_password=None if supplied_password else password))
        except (ValueError, HTTPException) as exc:
            db.rollback()
            results.append(BulkStudentResult(row=row_number, email=email or None, roll_no=roll_no or None, status="failed", error=str(exc.detail if isinstance(exc, HTTPException) else exc)))
    workbook.close()
    return BulkStudentImportOut(created=sum(result.status == "created" for result in results), failed=sum(result.status == "failed" for result in results), results=results)


async def _stage_import(db: Session, *, actor_id: UUID, kind: str, spreadsheet: UploadFile, images_zip: UploadFile | None = None, group_id: UUID | None = None) -> ImportJob:
    if not (spreadsheet.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="Upload an .xlsx or .xlsm Excel file")
    spreadsheet_bytes = await spreadsheet.read(settings.max_spreadsheet_bytes + 1)
    if len(spreadsheet_bytes) > settings.max_spreadsheet_bytes:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Spreadsheet is too large")
    spreadsheet_key = f"imports/pending/{secrets.token_urlsafe(24)}.xlsx"
    storage.put_bytes(spreadsheet_key, spreadsheet_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    zip_key = None
    try:
        if images_zip is not None:
            content = await images_zip.read(settings.max_bulk_zip_bytes + 1)
            if len(content) > settings.max_bulk_zip_bytes:
                raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="images_zip is too large")
            zip_key = f"imports/pending/{secrets.token_urlsafe(24)}.zip"
            storage.put_bytes(zip_key, content, "application/zip")
        job = ImportJob(kind=kind, created_by_id=actor_id, group_id=group_id, spreadsheet_key=spreadsheet_key, images_zip_key=zip_key)
        db.add(job)
        db.commit()
        db.refresh(job)
        process_import_job.delay(job.id)
        return job
    except Exception:
        storage.delete(spreadsheet_key)
        if zip_key:
            storage.delete(zip_key)
        raise


@router.post("/students/bulk", response_model=ImportJobOut, status_code=status.HTTP_202_ACCEPTED)
async def queue_bulk_student_creation(
    db: DbSession,
    spreadsheet: UploadFile = File(...),
    images_zip: UploadFile | None = File(None),
    user: User = Depends(require_roles(UserRole.admin)),
) -> ImportJobOut:
    return await _stage_import(db, actor_id=user.id, kind="student_create", spreadsheet=spreadsheet, images_zip=images_zip)


@router.post("/teachers/bulk", response_model=ImportJobOut, status_code=status.HTTP_202_ACCEPTED)
async def queue_bulk_teacher_creation(
    db: DbSession,
    spreadsheet: UploadFile = File(...),
    user: User = Depends(require_roles(UserRole.admin)),
) -> ImportJobOut:
    return await _stage_import(db, actor_id=user.id, kind="teacher_create", spreadsheet=spreadsheet)


@router.post("/students/bulk-update", response_model=ImportJobOut, status_code=status.HTTP_202_ACCEPTED)
async def queue_bulk_student_update(
    db: DbSession,
    spreadsheet: UploadFile = File(...),
    user: User = Depends(require_roles(UserRole.admin)),
) -> ImportJobOut:
    return await _stage_import(db, actor_id=user.id, kind="student_update", spreadsheet=spreadsheet)


@router.get("/photo-jobs/{job_id}", response_model=PhotoJobOut)
def photo_job_status(job_id: UUID, db: DbSession, _: User = Depends(require_roles(UserRole.admin))) -> PhotoProcessingJob:
    job = db.get(PhotoProcessingJob, job_id)
    if job is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Photo job not found")
    return job
