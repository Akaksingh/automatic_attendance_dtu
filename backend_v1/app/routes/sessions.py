"""Class-session lifecycle, classroom image review, and attendance correction."""
from __future__ import annotations

from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Annotated
from uuid import UUID, uuid4

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile, status
from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from ..config import settings
from ..database import get_db
from ..models import AuditLog, AttendanceRecord, AttendanceStatus, ClassSession, Group, GroupMembership, SessionPhoto, SessionState, Student, User, UserRole
from ..schemas import AttendanceStatusUpdate, PageMeta, SessionAttendanceOut, SessionOut, SessionPage, SessionPhotoOut
from ..security import CurrentUser, require_roles
from ..storage import storage
from ..tasks import expire_session_photo, process_class_session

router = APIRouter(prefix="/api", tags=["sessions"])
DbSession = Annotated[Session, Depends(get_db)]


def _get_session(session_id: UUID, db: Session) -> ClassSession:
    session = db.get(ClassSession, session_id)
    if session is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session not found")
    return session


def _staff_can_manage(session: ClassSession, user: User) -> bool:
    return user.role == UserRole.admin or (user.role == UserRole.teacher and session.group.owner_id == user.id)


def _authorized(session: ClassSession, user: User, db: Session) -> bool:
    if _staff_can_manage(session, user):
        return True
    return user.role == UserRole.student and db.query(GroupMembership).filter_by(group_id=session.group_id, student_id=user.id).first() is not None


def _photo_url(request: Request, photo_id: UUID) -> str:
    # Return a same-origin path so browser clients keep their authenticated
    # cookies when the API is behind Vite, nginx, or another reverse proxy.
    return str(request.url_for("session_photo_image", photo_id=photo_id).path)


def _to_out(session: ClassSession, request: Request, user: User, db: Session) -> SessionOut:
    attendance = session.attendance_records
    if user.role == UserRole.student:
        attendance = [record for record in attendance if record.student_id == user.id]
    return SessionOut(
        id=session.id,
        group_id=session.group_id,
        created_by_id=session.created_by_id,
        state=session.state,
        started_at=session.started_at,
        finalized_at=session.finalized_at,
        processing_total=session.processing_total,
        processing_completed=session.processing_completed,
        processing_message=session.processing_message,
        photos=[SessionPhotoOut(id=photo.id, image_url=_photo_url(request, photo.id), expires_at=photo.expires_at, processed_at=photo.processed_at, error=photo.error, created_at=photo.created_at) for photo in session.photos],
        attendance=[SessionAttendanceOut(id=record.id, student_id=record.student_id, status=record.status, marked_at=record.marked_at) for record in attendance],
    )


@router.get("/sessions/{session_id}/attendance.xlsx")
def export_session_attendance(session_id: UUID, db: DbSession, user: User = Depends(require_roles(UserRole.admin, UserRole.teacher))) -> Response:
    session = _get_session(session_id, db)
    if not _staff_can_manage(session, user):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You cannot export this session")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"
    sheet.append(["Roll no.", "Student", "Status", "Marked at"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
    rows = (
        db.query(Student, User, AttendanceRecord)
        .join(User, User.id == Student.user_id)
        .join(AttendanceRecord, AttendanceRecord.student_id == User.id)
        .filter(AttendanceRecord.session_id == session.id)
        .order_by(Student.roll_no.asc())
        .all()
    )
    for student, user_row, record in rows:
        sheet.append([student.roll_no, student.full_name, record.status.value.title(), record.marked_at.replace(tzinfo=None)])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 28
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["D"].width = 24
    output = BytesIO()
    workbook.save(output)
    filename = f"session_{session.started_at.strftime('%Y-%m-%d')}_attendance.xlsx"
    return Response(content=output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.post("/groups/{group_id}/sessions", response_model=SessionOut, status_code=status.HTTP_201_CREATED)
def create_session(
    group_id: UUID,
    request: Request,
    db: DbSession,
    user: User = Depends(require_roles(UserRole.admin, UserRole.teacher)),
) -> SessionOut:
    group = db.get(Group, group_id)
    if group is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Group not found")
    if group.archived_at is not None:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Archived groups cannot start sessions")
    if user.role != UserRole.admin and group.owner_id != user.id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You can only start sessions for groups you created")
    session = ClassSession(group_id=group.id, created_by_id=user.id, state=SessionState.created)
    db.add(session)
    db.commit()
    db.refresh(session)
    return _to_out(session, request, user, db)


@router.get("/groups/{group_id}/sessions", response_model=SessionPage)
def list_sessions(group_id: UUID, request: Request, db: DbSession, user: CurrentUser, page: int = Query(1, ge=1), page_size: int = Query(25, ge=1, le=100)) -> SessionPage:
    group = db.get(Group, group_id)
    if group is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Group not found")
    if user.role == UserRole.admin:
        pass
    elif user.role == UserRole.teacher and group.owner_id == user.id:
        pass
    elif user.role == UserRole.student and db.query(GroupMembership).filter_by(group_id=group_id, student_id=user.id).first():
        pass
    else:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You cannot access these sessions")
    query = db.query(ClassSession).filter(ClassSession.group_id == group_id).order_by(ClassSession.started_at.desc())
    total = query.count()
    sessions = query.offset((page - 1) * page_size).limit(page_size).all()
    return SessionPage(items=[_to_out(session, request, user, db) for session in sessions], meta=PageMeta(page=page, page_size=page_size, total=total, pages=(total + page_size - 1) // page_size))


@router.post("/sessions/{session_id}/photos", response_model=SessionOut, status_code=status.HTTP_202_ACCEPTED)
async def upload_session_photos(
    session_id: UUID,
    request: Request,
    db: DbSession,
    files: list[UploadFile] = File(...),
    user: User = Depends(require_roles(UserRole.admin, UserRole.teacher)),
) -> SessionOut:
    session = _get_session(session_id, db)
    if not _staff_can_manage(session, user):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You cannot manage this session")
    if session.state != SessionState.created:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Photos can only be added before processing starts")
    if not files:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Upload at least one classroom image")
    if len(session.photos) + len(files) > settings.max_session_photos:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail=f"A session can contain at most {settings.max_session_photos} classroom images")
    staged: list[str] = []
    try:
        expires_at = datetime.now(UTC) + timedelta(days=settings.session_photo_retention_days)
        for file in files:
            if file.content_type not in {"image/jpeg", "image/png", "image/webp"}:
                raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="Use JPEG, PNG, or WebP classroom images")
            content = await file.read()
            if len(content) > settings.max_profile_photo_bytes:
                raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Classroom image is too large")
            key = f"{session.id}/{uuid4()}{Path(file.filename or '').suffix.lower() or '.jpg'}"
            key = f"sessions/{key}"
            storage.put_bytes(key, content, file.content_type)
            staged.append(key)
            db.add(SessionPhoto(session_id=session.id, storage_key=key, content_type=file.content_type, expires_at=expires_at))
        db.commit()
    except Exception:
        db.rollback()
        for key in staged:
            storage.delete(key)
        raise
    db.refresh(session)
    return _to_out(session, request, user, db)


@router.post("/sessions/{session_id}/process", response_model=SessionOut, status_code=status.HTTP_202_ACCEPTED)
def process_session(
    session_id: UUID,
    request: Request,
    db: DbSession,
    user: User = Depends(require_roles(UserRole.admin, UserRole.teacher)),
) -> SessionOut:
    session = _get_session(session_id, db)
    if not _staff_can_manage(session, user):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You cannot manage this session")
    if session.state != SessionState.created:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="This session is already processing or has been reviewed")
    if not session.photos:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Add at least one classroom image before processing")
    session.state = SessionState.processing
    db.commit()
    db.refresh(session)
    process_class_session.delay(session.id)
    for photo in session.photos:
        expire_session_photo.apply_async(args=[photo.id], countdown=settings.session_photo_retention_days * 86400)
    return _to_out(session, request, user, db)


@router.get("/sessions/{session_id}/progress", response_model=SessionOut)
def session_progress(session_id: UUID, request: Request, db: DbSession, user: CurrentUser) -> SessionOut:
    """Polling endpoint for worker progress; the same data is also in GET /sessions/{id}."""
    session = _get_session(session_id, db)
    if not _authorized(session, user, db):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You cannot access this session")
    return _to_out(session, request, user, db)


@router.get("/sessions/{session_id}", response_model=SessionOut)
def get_session(session_id: UUID, request: Request, db: DbSession, user: CurrentUser) -> SessionOut:
    session = _get_session(session_id, db)
    if not _authorized(session, user, db):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You cannot access this session")
    return _to_out(session, request, user, db)


@router.get("/sessions/photos/{photo_id}/image", name="session_photo_image", include_in_schema=False)
def session_photo_image(photo_id: UUID, user: CurrentUser, db: DbSession):
    photo = db.get(SessionPhoto, photo_id)
    if photo is None or photo.expires_at <= datetime.now(UTC):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session image is unavailable")
    session = _get_session(photo.session_id, db)
    if not _authorized(session, user, db):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You cannot access this image")
    if not storage.exists(photo.storage_key):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session image is unavailable")
    return StreamingResponse(BytesIO(storage.get_bytes(photo.storage_key)), media_type=photo.content_type)


@router.patch("/sessions/{session_id}/attendance/{student_id}", response_model=SessionAttendanceOut)
def update_attendance(
    session_id: UUID,
    student_id: UUID,
    payload: AttendanceStatusUpdate,
    db: DbSession,
    user: User = Depends(require_roles(UserRole.admin, UserRole.teacher)),
) -> AttendanceRecord:
    session = _get_session(session_id, db)
    if not _staff_can_manage(session, user):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You cannot modify this session")
    if session.state == SessionState.finalized:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Finalized attendance cannot be changed")
    if not db.query(GroupMembership).filter_by(group_id=session.group_id, student_id=student_id).first():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Student is not in this group")
    record = db.query(AttendanceRecord).filter_by(session_id=session.id, student_id=student_id).first()
    if record is None:
        record = AttendanceRecord(session_id=session.id, student_id=student_id, status=payload.status)
        db.add(record)
    else:
        record.status = payload.status
        record.marked_at = datetime.now(UTC)
    db.add(AuditLog(actor_id=user.id, action="attendance.updated", entity_type="attendance_record", entity_id=record.id))
    db.commit()
    db.refresh(record)
    return record


@router.post("/sessions/{session_id}/finalize", response_model=SessionOut)
def finalize_session(session_id: UUID, request: Request, db: DbSession, user: User = Depends(require_roles(UserRole.admin, UserRole.teacher))) -> SessionOut:
    session = _get_session(session_id, db)
    if not _staff_can_manage(session, user):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You cannot finalize this session")
    if session.state != SessionState.review:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Session must finish processing before finalization")
    session.state = SessionState.finalized
    session.finalized_at = datetime.now(UTC)
    db.add(AuditLog(actor_id=user.id, action="session.finalized", entity_type="class_session", entity_id=session.id))
    db.commit()
    db.refresh(session)
    return _to_out(session, request, user, db)


@router.delete("/sessions/{session_id}", status_code=status.HTTP_204_NO_CONTENT, response_class=Response)
def delete_session(session_id: UUID, db: DbSession, user: User = Depends(require_roles(UserRole.admin, UserRole.teacher))) -> Response:
    session = _get_session(session_id, db)
    if not _staff_can_manage(session, user):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You cannot delete this session")
    for photo in session.photos:
        storage.delete(photo.storage_key)
    db.add(AuditLog(actor_id=user.id, action="session.deleted", entity_type="class_session", entity_id=session.id))
    db.delete(session)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)
